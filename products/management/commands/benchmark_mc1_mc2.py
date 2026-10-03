import os
import time
import random
import string
from dataclasses import dataclass
from datetime import datetime, timedelta
from django.core.management.base import BaseCommand
from engine.datastructures.hash_table import HashTable
from engine.datastructures.max_heap import MaxHeap
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


@dataclass
class MockProduct:
    id: int
    sku: str
    name: str
    price: float


@dataclass
class MockOrder:
    id: int
    order_code: str
    priority: int
    created_at: datetime


class NaiveProductLookup:
    """Tuyen tinh cho MC1"""
    def __init__(self):
        self.items = []

    def bulk_load(self, products):
        self.items = list(products)

    def lookup(self, sku):
        for item in self.items:
            if item.sku == sku:
                return item
        return None


class NaiveOrderQueue:
    """Tuyen tinh cho MC2"""
    def __init__(self):
        self.items = []

    def bulk_load(self, orders):
        self.items = list(orders)

    def pop_max(self):
        if not self.items:
            return None
        best_idx = 0
        best_key = (
            self.items[0].priority,
            -self.items[0].created_at.timestamp() if self.items[0].created_at else 0,
            -self.items[0].id
        )
        for i in range(1, len(self.items)):
            item = self.items[i]
            item_key = (
                item.priority,
                -item.created_at.timestamp() if item.created_at else 0,
                -item.id
            )
            if item_key > best_key:
                best_key = item_key
                best_idx = i
        return self.items.pop(best_idx)


class Command(BaseCommand):
    help = "Benchmark MC1 (HashTable) and MC2 (MaxHeap) against naive baselines."

    def add_arguments(self, parser):
        parser.add_argument(
            '--scales',
            type=str,
            default='1000,10000,100000',
            help='List of N scales separated by commas (default: 1000,10000,100000)'
        )
        parser.add_argument(
            '--lookup-queries',
            type=int,
            default=1000,
            help='Number of lookup queries for MC1 (default: 1000)'
        )
        parser.add_argument(
            '--dispatch-ops',
            type=int,
            default=500,
            help='Number of pop_max operations for MC2 (default: 500)'
        )
        parser.add_argument(
            '--output',
            type=str,
            default='benchmark_results.xlsx',
            help='Excel output file path (default: benchmark_results.xlsx)'
        )

    def handle(self, *args, **options):
        scales = [int(s.strip()) for s in options['scales'].split(',') if s.strip()]
        lookup_queries_count = options['lookup_queries']
        dispatch_ops_count = options['dispatch_ops']
        output_file = options['output']

        self.stdout.write(f"=== Benchmark MC1 & MC2 ===")
        self.stdout.write(f"Scales: {scales} | Lookups (K): {lookup_queries_count:,} | Dispatches (K): {dispatch_ops_count:,}\n")

        # ---------------------------------------------------------
        # 1. MC1: Product Lookup
        # ---------------------------------------------------------
        self.stdout.write("[MC1] Product Lookup (HashTable vs Linear Search)")
        mc1_results = []
        for n in scales:
            products = []
            skus = []
            for i in range(n):
                sku = f"SKU-{i:07d}-{''.join(random.choices(string.ascii_uppercase, k=3))}"
                skus.append(sku)
                products.append(MockProduct(id=i + 1, sku=sku, name=f"Product {i + 1}", price=100000))

            query_skus = random.sample(skus, min(lookup_queries_count, n))
            if len(query_skus) < lookup_queries_count:
                query_skus += [f"NONEXIST-{i}" for i in range(lookup_queries_count - len(query_skus))]

            # HashTable Bulk Load
            ht = HashTable(size=max(1000, n // 2))
            t0 = time.perf_counter()
            for p in products:
                ht.put(p.sku, p)
            ht_load_time = (time.perf_counter() - t0) * 1000

            # HashTable Lookup
            t0 = time.perf_counter()
            for q in query_skus:
                _ = ht.get(q)
            ht_lookup_time = (time.perf_counter() - t0) * 1000

            # Naive Bulk Load
            naive_lookup = NaiveProductLookup()
            t0 = time.perf_counter()
            naive_lookup.bulk_load(products)
            naive_load_time = (time.perf_counter() - t0) * 1000

            # Naive Lookup
            t0 = time.perf_counter()
            for q in query_skus:
                _ = naive_lookup.lookup(q)
            naive_lookup_time = (time.perf_counter() - t0) * 1000

            speedup = naive_lookup_time / ht_lookup_time if ht_lookup_time > 0 else 0
            mc1_results.append({
                'n': n,
                'ht_load': ht_load_time,
                'ht_lookup': ht_lookup_time,
                'naive_load': naive_load_time,
                'naive_lookup': naive_lookup_time,
                'speedup': speedup
            })

        self.stdout.write(f"{'N':<10} | {'HT Load (ms)':<14} | {'HT Lookup (ms)':<15} | {'Naive Load (ms)':<16} | {'Naive Lookup (ms)':<18} | {'Speedup':<10}")
        self.stdout.write("-" * 92)
        for r in mc1_results:
            self.stdout.write(
                f"{r['n']:<10,}"
                f" | {r['ht_load']:>12.2f}"
                f" | {r['ht_lookup']:>13.2f}"
                f" | {r['naive_load']:>14.2f}"
                f" | {r['naive_lookup']:>16.2f}"
                f" | {r['speedup']:>8.1f}x"
            )

        # ---------------------------------------------------------
        # 2. MC2: Order Queue
        # ---------------------------------------------------------
        self.stdout.write("\n[MC2] Order Queue Dispatch (MaxHeap vs Unsorted List)")
        mc2_results = []
        now = datetime.now()
        for n in scales:
            orders = []
            for i in range(n):
                orders.append(MockOrder(
                    id=i + 1,
                    order_code=f"ORD-{i:07d}",
                    priority=random.randint(1, 4),
                    created_at=now - timedelta(minutes=random.randint(1, 10000))
                ))

            k_ops = min(dispatch_ops_count, n)

            # MaxHeap build_heap
            heap = MaxHeap(key_func=lambda o: (o.priority, -o.created_at.timestamp() if o.created_at else 0, -o.id))
            t0 = time.perf_counter()
            heap.build_heap(orders)
            heap_build_time = (time.perf_counter() - t0) * 1000

            # MaxHeap Pop K
            t0 = time.perf_counter()
            for _ in range(k_ops):
                _ = heap.pop_max()
            heap_pop_time = (time.perf_counter() - t0) * 1000

            # Naive Bulk Load
            naive_queue = NaiveOrderQueue()
            t0 = time.perf_counter()
            naive_queue.bulk_load(orders)
            naive_load_time = (time.perf_counter() - t0) * 1000

            # Naive Pop K
            t0 = time.perf_counter()
            for _ in range(k_ops):
                _ = naive_queue.pop_max()
            naive_pop_time = (time.perf_counter() - t0) * 1000

            speedup = naive_pop_time / heap_pop_time if heap_pop_time > 0 else 0
            mc2_results.append({
                'n': n,
                'heap_build': heap_build_time,
                'heap_pop': heap_pop_time,
                'naive_load': naive_load_time,
                'naive_pop': naive_pop_time,
                'speedup': speedup,
                'k': k_ops
            })

        self.stdout.write(f"{'N':<10} | {'Heap Build (ms)':<15} | {'Heap Pop (ms)':<15} | {'Naive Load (ms)':<16} | {'Naive Pop (ms)':<16} | {'Speedup':<10}")
        self.stdout.write("-" * 92)
        for r in mc2_results:
            self.stdout.write(
                f"{r['n']:<10,}"
                f" | {r['heap_build']:>13.2f}"
                f" | {r['heap_pop']:>13.2f}"
                f" | {r['naive_load']:>14.2f}"
                f" | {r['naive_pop']:>14.2f}"
                f" | {r['speedup']:>8.1f}x"
            )

        # ---------------------------------------------------------
        # 3. Export to Excel
        # ---------------------------------------------------------
        saved_path = self.export_excel(output_file, mc1_results, mc2_results, lookup_queries_count, dispatch_ops_count)
        self.stdout.write(f"\n[OK] Excel file saved: {saved_path}")

    def export_excel(self, file_path, mc1_results, mc2_results, k_mc1, k_mc2):
        wb = openpyxl.Workbook()

        # Styles
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        data_font = Font(name="Segoe UI", size=10)
        bold_font = Font(name="Segoe UI", size=10, bold=True)
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        center_align = Alignment(horizontal="center", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")

        # --- Sheet 1: MC1 ---
        ws1 = wb.active
        ws1.title = "MC1_Product_Lookup"

        headers_mc1 = [
            "Quy mô (N)",
            "HashTable Bulk Load (ms)",
            f"HashTable Tra cứu K={k_mc1:,} (ms)",
            "Naive Bulk Load (ms)",
            f"Naive Tra cứu K={k_mc1:,} (ms)",
            "Tốc độ tăng tốc (Speedup)"
        ]

        ws1.append(headers_mc1)
        for col_num, cell in enumerate(ws1[1], 1):
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        for row_idx, r in enumerate(mc1_results, start=2):
            ws1.append([
                r['n'],
                round(r['ht_load'], 2),
                round(r['ht_lookup'], 2),
                round(r['naive_load'], 2),
                round(r['naive_lookup'], 2),
                round(r['speedup'], 1)
            ])
            for col_idx in range(1, 7):
                cell = ws1.cell(row=row_idx, column=col_idx)
                cell.font = bold_font if col_idx in (1, 6) else data_font
                cell.border = thin_border
                cell.alignment = center_align if col_idx == 1 else right_align
                if col_idx == 1:
                    cell.number_format = '#,##0'
                elif col_idx == 6:
                    cell.number_format = '0.0"x"'
                else:
                    cell.number_format = '0.00'

        self._auto_fit_columns(ws1)

        # --- Sheet 2: MC2 ---
        ws2 = wb.create_sheet(title="MC2_Order_Queue")
        headers_mc2 = [
            "Quy mô (N)",
            "MaxHeap build_heap (ms)",
            f"MaxHeap Lấy K={k_mc2:,} đơn (ms)",
            "Naive Bulk Load (ms)",
            f"Naive Lấy K={k_mc2:,} đơn (ms)",
            "Tốc độ tăng tốc (Speedup)"
        ]

        ws2.append(headers_mc2)
        for col_num, cell in enumerate(ws2[1], 1):
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        for row_idx, r in enumerate(mc2_results, start=2):
            ws2.append([
                r['n'],
                round(r['heap_build'], 2),
                round(r['heap_pop'], 2),
                round(r['naive_load'], 2),
                round(r['naive_pop'], 2),
                round(r['speedup'], 1)
            ])
            for col_idx in range(1, 7):
                cell = ws2.cell(row=row_idx, column=col_idx)
                cell.font = bold_font if col_idx in (1, 6) else data_font
                cell.border = thin_border
                cell.alignment = center_align if col_idx == 1 else right_align
                if col_idx == 1:
                    cell.number_format = '#,##0'
                elif col_idx == 6:
                    cell.number_format = '0.0"x"'
                else:
                    cell.number_format = '0.00'

        self._auto_fit_columns(ws2)

        # Save workbook with fallback if file is currently open / locked
        try:
            wb.save(file_path)
            return file_path
        except PermissionError:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            base, ext = os.path.splitext(file_path)
            fallback_path = f"{base}_{timestamp}{ext or '.xlsx'}"
            self.stdout.write(self.style.WARNING(f"\n[Warning] File '{file_path}' is open in another program."))
            self.stdout.write(self.style.WARNING(f"          Saved fallback to: '{fallback_path}'"))
            wb.save(fallback_path)
            return fallback_path

    def _auto_fit_columns(self, ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)
